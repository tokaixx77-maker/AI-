import streamlit as st
import pandas as pd
import numpy as np
import random
import json
import os
import tempfile
import time
import io
from http import HTTPStatus
from concurrent.futures import ThreadPoolExecutor, as_completed

# 引入 Excel 格式处理库
try:
    import openpyxl
    from openpyxl.styles import Font, Alignment
    import dashscope
except ImportError:
    st.error("❌ 缺少必要库，请在终端运行: pip install dashscope openpyxl streamlit pandas")
    st.stop()

# ==========================================
# 🎨 UI & 专属 Logo (液态金属风格 + 醒目按钮)
# ==========================================
st.set_page_config(page_title="77智能体测系统", layout="wide", page_icon="⚡")

def render_ui_header():
    st.markdown("""
    <style>
    @import url('https://fonts.googleapis.com/css2?family=Montserrat:wght@900&display=swap');
    
    /* 主LOGO样式 */
    .main-logo {
        font-family: 'Montserrat', 'Arial Black', sans-serif;
        font-size: 80px;
        font-weight: 900;
        background: linear-gradient(135deg, #e0e0e0 0%, #ffffff 25%, #8a94a3 50%, #ffffff 75%, #c0c5d0 100%);
        background-size: 150% auto;
        -webkit-background-clip: text;
        -webkit-text-fill-color: transparent;
        text-shadow: 1px 1px 1px rgba(255,255,255,0.8), 0px 5px 15px rgba(0,0,0,0.3);
        margin-bottom: -15px;
    }
    
    .sub-title { font-family: '微软雅黑', sans-serif; color: #666; font-size: 16px; letter-spacing: 1px; font-weight: 600; }
    
    /* 进度条颜色 */
    .stProgress > div > div > div > div { background-color: #5a6e8c; }
    
    /* 醒目按钮样式重写 */
    div.stButton > button {
        width: 100%;
        border-radius: 8px;
        font-weight: bold;
        transition: all 0.3s ease;
    }
    
    /* 主要操作按钮 (Primary) */
    div.stButton > button[kind="primary"] {
        background: linear-gradient(90deg, #0575E6 0%, #021B79 100%);
        border: none;
        box-shadow: 0 4px 15px rgba(5, 117, 230, 0.4);
    }
    div.stButton > button[kind="primary"]:hover {
        transform: translateY(-2px);
        box-shadow: 0 6px 20px rgba(5, 117, 230, 0.6);
    }

    /* 成功/确认类提示框 */
    .success-box {
        padding: 1rem;
        background-color: #d4edda;
        color: #155724;
        border-radius: 0.5rem;
        border-left: 5px solid #28a745;
        margin-bottom: 1rem;
    }
    
    /* 仪表盘样式 */
    .metric-card {
        background-color: #f8f9fa;
        border: 1px solid #e9ecef;
        padding: 15px;
        border-radius: 10px;
        text-align: center;
        box-shadow: 0 2px 5px rgba(0,0,0,0.05);
    }
    </style>
    <div>
        <h1 class='main-logo'>77 SYSTEM</h1>
        <p class='sub-title'>全学段体育数据智能中枢 // v6.6 Final-Audit</p>
    </div>
    <hr style="border: 0; height: 1px; background-image: linear-gradient(to right, rgba(0,0,0,0), rgba(0,0,0,0.2), rgba(0,0,0,0));">
    """, unsafe_allow_html=True)

# ==========================================
# ⚙️ 核心配置
# ==========================================
# === API Key 安全加载 (修改后) ===
import os
try:
    # 尝试从 Streamlit 云端的“保险箱”(Secrets)读取
    MY_API_KEY = st.secrets["DASHSCOPE_API_KEY"]
except:
    # 如果在本地没配置保险箱，尝试读取环境变量，或者留空
    MY_API_KEY = os.getenv("DASHSCOPE_API_KEY", "")
SCORE_COLUMNS_MAP = {
    '50米': ['50米', '50米跑', '五十米'],
    '立定跳远': ['立定跳远', '跳远'],
    '坐位体前屈': ['坐位体前屈', '体前屈', '坐位'],
    '1000米': ['1000米', '1000米跑', '一千米'],
    '800米': ['800米', '800米跑', '八百米'],
    '引体向上': ['引体向上', '引体'],
    '仰卧起坐': ['仰卧起坐', '一分钟仰卧起坐'],
    '肺活量': ['肺活量'],
    '身高': ['身高'],
    '体重': ['体重'],
    '跳绳': ['跳绳', '一分钟跳绳'],
    '50x8往返跑': ['50x8', '50米x8', '往返跑']
}

# ==========================================
# 📚 核心功能: 加载标准库 (更完整的阈值数据)
# ==========================================
# 这里包含了各个学段主要项目的及格(60)、良好(80)、优秀(90)参考线
# 用于数据调整算法判断
DEFAULT_STANDARDS_DATA = {
  "小学(低年级)": {
    "说明": "适用于小学1-2年级",
    "男": {
      "50米": {"60": 12.0, "80": 10.6, "90": 10.2},
      "坐位体前屈": {"60": -0.6, "80": 6.9, "90": 10.1},
      "立定跳远": {"60": 117, "80": 140, "90": 149},
      "跳绳": {"60": 17, "80": 80, "90": 99},
      "肺活量": {"60": 680, "80": 1080, "90": 1250}
    },
    "女": {
      "50米": {"60": 12.6, "80": 11.2, "90": 10.7},
      "坐位体前屈": {"60": 1.6, "80": 8.7, "90": 11.5},
      "立定跳远": {"60": 110, "80": 132, "90": 139},
      "跳绳": {"60": 17, "80": 84, "90": 103},
      "肺活量": {"60": 580, "80": 900, "90": 1050}
    }
  },
  "小学(中年级)": {
    "说明": "适用于小学3-4年级",
    "男": {
      "50米": {"60": 10.6, "80": 9.2, "90": 8.8},
      "坐位体前屈": {"60": 0.5, "80": 8.0, "90": 11.3},
      "立定跳远": {"60": 138, "80": 161, "90": 174},
      "跳绳": {"60": 40, "80": 95, "90": 116},
      "仰卧起坐": {"60": 23, "80": 39, "90": 45},
      "肺活量": {"60": 1100, "80": 1700, "90": 1950}
    },
    "女": {
      "50米": {"60": 10.8, "80": 9.5, "90": 9.1},
      "坐位体前屈": {"60": 2.8, "80": 9.8, "90": 12.9},
      "立定跳远": {"60": 130, "80": 149, "90": 158},
      "跳绳": {"60": 42, "80": 99, "90": 121},
      "仰卧起坐": {"60": 21, "80": 37, "90": 43},
      "肺活量": {"60": 950, "80": 1350, "90": 1550}
    }
  },
  "小学(高年级)": {
    "说明": "适用于小学5-6年级",
    "男": {
      "50米": {"60": 9.8, "80": 8.6, "90": 8.2},
      "坐位体前屈": {"60": 2.0, "80": 9.0, "90": 12.5},
      "立定跳远": {"60": 161, "80": 193, "90": 205},
      "跳绳": {"60": 65, "80": 126, "90": 148},
      "仰卧起坐": {"60": 28, "80": 42, "90": 48},
      "50x8往返跑": {"60": 138, "80": 122, "90": 116},
      "肺活量": {"60": 1400, "80": 2200, "90": 2500}
    },
    "女": {
      "50米": {"60": 10.2, "80": 9.1, "90": 8.7},
      "坐位体前屈": {"60": 4.5, "80": 10.5, "90": 13.5},
      "立定跳远": {"60": 139, "80": 164, "90": 175},
      "跳绳": {"60": 67, "80": 129, "90": 154},
      "仰卧起坐": {"60": 27, "80": 41, "90": 47},
      "50x8往返跑": {"60": 144, "80": 128, "90": 122},
      "肺活量": {"60": 1150, "80": 1650, "90": 1850}
    }
  },
  "初中": {
    "说明": "适用于初中7-9年级",
    "男": {
      "50米": {"60": 9.1, "80": 7.9, "90": 7.3},
      "坐位体前屈": {"60": 3.7, "80": 12.8, "90": 16.6},
      "立定跳远": {"60": 185, "80": 225, "90": 240},
      "引体向上": {"60": 6, "80": 11, "90": 14},
      "1000米": {"60": 275, "80": 245, "90": 225},
      "肺活量": {"60": 2300, "80": 3400, "90": 3900}
    },
    "女": {
      "50米": {"60": 9.5, "80": 8.5, "90": 7.9},
      "坐位体前屈": {"60": 6.0, "80": 14.8, "90": 18.2},
      "立定跳远": {"60": 146, "80": 176, "90": 190},
      "仰卧起坐": {"60": 28, "80": 40, "90": 46},
      "800米": {"60": 255, "80": 225, "90": 205},
      "肺活量": {"60": 1800, "80": 2500, "90": 2800}
    }
  },
  "高中": {
    "说明": "适用于高中10-12年级",
    "男": {
      "50米": {"60": 9.0, "80": 7.8, "90": 7.2},
      "坐位体前屈": {"60": 5.5, "80": 14.9, "90": 18.8},
      "立定跳远": {"60": 205, "80": 240, "90": 252},
      "引体向上": {"60": 9, "80": 14, "90": 17},
      "1000米": {"60": 265, "80": 235, "90": 215},
      "肺活量": {"60": 2700, "80": 3800, "90": 4300}
    },
    "女": {
      "50米": {"60": 9.4, "80": 8.7, "90": 8.0},
      "坐位体前屈": {"60": 6.8, "80": 16.5, "90": 19.9},
      "立定跳远": {"60": 158, "80": 182, "90": 194},
      "仰卧起坐": {"60": 28, "80": 40, "90": 46},
      "800米": {"60": 270, "80": 240, "90": 218},
      "肺活量": {"60": 1850, "80": 2650, "90": 2950}
    }
  },
  "大学": {
    "说明": "适用于大学1-4年级",
    "男": {
      "50米": {"60": 9.1, "80": 7.9, "90": 7.1},
      "坐位体前屈": {"60": 3.7, "80": 14.7, "90": 18.2},
      "立定跳远": {"60": 208, "80": 248, "90": 263},
      "引体向上": {"60": 10, "80": 15, "90": 17},
      "1000米": {"60": 272, "80": 237, "90": 212},
      "肺活量": {"60": 3100, "80": 4300, "90": 4800}
    },
    "女": {
      "50米": {"60": 10.3, "80": 8.9, "90": 8.3},
      "坐位体前屈": {"60": 6.0, "80": 16.9, "90": 20.3},
      "立定跳远": {"60": 151, "80": 178, "90": 195},
      "仰卧起坐": {"60": 26, "80": 42, "90": 50},
      "800米": {"60": 274, "80": 244, "90": 214},
      "肺活量": {"60": 2000, "80": 3000, "90": 3250}
    }
  }
}

def load_standards():
    """
    读取 standards.json 文件。如果不存在，则自动创建。
    """
    file_path = 'standards.json'
    if not os.path.exists(file_path):
        try:
            with open(file_path, 'w', encoding='utf-8') as f:
                json.dump(DEFAULT_STANDARDS_DATA, f, ensure_ascii=False, indent=2)
            st.toast("✅ 已自动生成标准库文件 (standards.json)", icon="🛠️")
        except Exception as e:
            st.error(f"❌ 无法创建标准库文件: {e}")
            return None
    try:
        with open(file_path, 'r', encoding='utf-8') as f:
            return json.load(f)
    except FileNotFoundError:
        return None

STANDARDS_DB = load_standards()

# ==========================================
# 📊 核心功能: 智能调整算法 (基于数据库)
# ==========================================
def calculate_good_rate(df, level):
    """计算当前数据的平均优良率（单项指标达到80分以上的比例）"""
    if not STANDARDS_DB or level not in STANDARDS_DB: return 0.0
    
    std_male = STANDARDS_DB[level]["男"]
    std_female = STANDARDS_DB[level]["女"]
    
    total_items = 0
    total_good = 0
    
    if len(df) == 0: return 0.0
    
    for _, row in df.iterrows():
        gender = row.get('性别')
        current_std = std_male if str(gender) in ['1','1.0','男'] else std_female
        
        for col in df.columns:
            key = None
            for k in current_std.keys():
                if k in col:
                    key = k
                    break
            
            if key and pd.notna(row[col]):
                try:
                    val = float(row[col])
                    good_cutoff = float(current_std[key]['80'])
                    is_running = key in ['50米', '800米', '1000米', '50x8往返跑', '50米跑', '800米跑', '1000米跑']
                    
                    if is_running:
                        if val <= good_cutoff: total_good += 1
                    else:
                        if val >= good_cutoff: total_good += 1
                    total_items += 1
                except:
                    pass
                
    if total_items == 0: return 0.0
    return (total_good / total_items) * 100

def auto_boost(df, target_rate, level):
    df_a = df.copy()
    if not STANDARDS_DB or level not in STANDARDS_DB:
        st.error(f"❌ 无法获取 [{level}] 的评分标准，请确认 standards.json 文件存在且完整。")
        return df_a

    std_male = STANDARDS_DB[level]["男"]
    std_female = STANDARDS_DB[level]["女"]

    # 1. 收集所有可调整的数据点及其状态
    modifiable_cells = [] 
    
    total_valid_items = 0
    current_good_items = 0

    # 遍历数据建立索引
    for idx, row in df_a.iterrows():
        gender = row.get('性别')
        current_std = std_male if str(gender) in ['1','1.0','男'] else std_female
        
        for col in df_a.columns:
            key = None
            for k in current_std.keys():
                if k in col:
                    key = k
                    break
            
            if key and pd.notna(row[col]):
                try:
                    val = float(row[col])
                    good_cutoff = float(current_std[key]['80'])
                    is_running = key in ['50米', '800米', '1000米', '50x8往返跑', '50米跑', '800米跑', '1000米跑']
                    
                    is_good = False
                    if is_running:
                        if val <= good_cutoff: is_good = True
                    else:
                        if val >= good_cutoff: is_good = True
                    
                    if is_good:
                        current_good_items += 1
                    
                    modifiable_cells.append({
                        'idx': idx,
                        'col': col,
                        'val': val,
                        'threshold': good_cutoff,
                        'is_good': is_good,
                        'is_running': is_running
                    })
                    total_valid_items += 1
                except:
                    continue

    if total_valid_items == 0:
        return df_a

    # 2. 计算缺口
    current_rate = (current_good_items / total_valid_items) * 100
    target_good_count = int(total_valid_items * (target_rate / 100))
    needed = target_good_count - current_good_items

    with st.status(f"🚀 正在执行智能算法 (目标: {target_rate}%)", expanded=True) as status:
        st.write(f"📊 当前优良率: {current_rate:.2f}%")
        
        if needed <= 0:
            st.success("✅ 当前数据已达到或超过目标优良率，无需调整。")
            time.sleep(1)
            status.update(label="处理完毕 (无需调整)", state="complete", expanded=False)
            return df_a
        
        st.write(f"⚡ 需要优化 {needed} 个数据点以达标...")
        progress_bar = st.progress(0)
        
        # 3. 筛选出非优良的项目，准备提升
        not_good_cells = [c for c in modifiable_cells if not c['is_good']]
        
        # 随机抽取需要提升的数量
        if len(not_good_cells) > needed:
            cells_to_boost = random.sample(not_good_cells, needed)
        else:
            cells_to_boost = not_good_cells 
            
        # 4. 执行调整
        for i, cell in enumerate(cells_to_boost):
            row_idx = cell['idx']
            col_name = cell['col']
            threshold = cell['threshold']
            is_running = cell['is_running']
            
            # 生成新值：确保肯定过线（真实感模拟）
            if is_running:
                # 跑步：比阈值少(快)一点，幅度在 0.05s 到 1.5s 之间，模拟自然分布
                # 防止生成负数
                boost_amount = random.uniform(0.05, 1.5)
                new_val = round(max(0.1, threshold - boost_amount), 2)
            else:
                # 数值：比阈值多一点
                if any(x in col_name for x in ['跳绳', '仰卧', '引体']):
                     # 次数类：多做 1-5 个
                     new_val = int(threshold + random.randint(1, 5))
                elif any(x in col_name for x in ['跳远', '身高', '体重']):
                     # 长度/重量：多 1-10 个单位 (cm/kg)
                     new_val = int(threshold + random.randint(1, 10))
                elif '肺活量' in col_name:
                     # 肺活量：多 50-300 ml
                     new_val = int(threshold + random.randint(50, 300))
                else:
                     # 其他（如坐位体前屈）：多 0.2-2.0
                     new_val = round(threshold + random.uniform(0.2, 2.0), 1)
            
            df_a.at[row_idx, col_name] = new_val
            
            if i % 10 == 0: progress_bar.progress((i + 1) / len(cells_to_boost))
            
        progress_bar.progress(100)
        st.write(f"✨ 已精准优化 {len(cells_to_boost)} 个数据点！")
        status.update(label="✅ 处理完毕", state="complete", expanded=False)

    return df_a

# ==========================================
# 🛠️ 辅助功能 (Excel回写, OCR, 清洗)
# ==========================================
def save_data_keeping_format(original_file_obj, processed_df):
    try:
        wb = openpyxl.load_workbook(original_file_obj)
        ws = wb.active
        header_map = {}
        for col_idx, cell in enumerate(ws[1], start=1):
            if cell.value: header_map[str(cell.value).strip()] = col_idx
        name_col_idx = header_map.get('姓名')
        if not name_col_idx: return None, "❌ 错误：在模板中未找到【姓名】列"
        name_row_map = {}
        for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=False), start=2):
            cell_val = row[name_col_idx-1].value
            if cell_val: name_row_map[str(cell_val).strip()] = row_idx
        
        for i, (idx, row) in enumerate(processed_df.iterrows()):
            name = str(row.get('姓名', '')).strip()
            target_row = name_row_map.get(name)
            if target_row:
                for col_name in processed_df.columns:
                    target_col = header_map.get(col_name)
                    if target_col and pd.notna(row[col_name]):
                        ws.cell(row=target_row, column=target_col).value = row[col_name]
        
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output, "✅ 格式无损导出准备就绪！"
    except Exception as e:
        return None, f"导出失败: {str(e)}"

def call_qwen_vl_ocr(image_file, api_key):
    dashscope.api_key = api_key
    prompt = """
    分析体测成绩单。提取：姓名、性别、班级、所有体育项目成绩。
    规则：
    1. 姓名必须准确。
    2. 立定跳远单位转厘米(2.3 -> 230)。
    3. 输出纯 JSON 列表。
    """
    # 处理文件流
    try:
        # 重置指针以防万一
        image_file.seek(0)
        tfile = tempfile.NamedTemporaryFile(delete=False, suffix=image_file.name) 
        tfile.write(image_file.read())
        local_path = tfile.name
        tfile.close()
        
        messages = [{"role": "user", "content": [{"image": f"file://{local_path}"}, {"text": prompt}]}]
        response = dashscope.MultiModalConversation.call(model='qwen-vl-max', messages=messages)
        
        # 清理临时文件
        os.remove(local_path)
        
        if response.status_code == HTTPStatus.OK:
            content = response.output.choices[0].message.content[0]['text']
            return content.replace("```json", "").replace("```", "").strip()
        else:
            return f"Error: {response.code} - {response.message}"
    except Exception as e:
        if 'local_path' in locals() and os.path.exists(local_path): os.remove(local_path)
        return f"API Error: {str(e)}"

def merge_ocr_to_master(master_df, ocr_df):
    merged_df = master_df.copy()
    logs = []
    
    # 记录详细变更
    changes_list = []
    
    ocr_cols = ocr_df.columns
    col_mapping = {}
    for std_col, alias_list in SCORE_COLUMNS_MAP.items():
        for alias in alias_list:
            if alias in ocr_cols: col_mapping[alias] = std_col; break
    ocr_df = ocr_df.rename(columns=col_mapping)
    
    with st.spinner("🔄 正在智能匹配合并..."):
        time.sleep(0.5)
        
    for idx, row in ocr_df.iterrows():
        name = str(row.get('姓名', '')).strip()
        if not name: continue
        
        match_mask = merged_df['姓名'].astype(str).str.strip() == name
        if match_mask.any():
            target_idx = merged_df[match_mask].index[0]
            items_updated = []
            for col in row.index:
                if col in SCORE_COLUMNS_MAP.keys() and col in merged_df.columns:
                    val = row[col]
                    if pd.notna(val): 
                        old_val = merged_df.at[target_idx, col]
                        merged_df.at[target_idx, col] = val
                        items_updated.append(col)
                        
                        # 记录变更详情
                        changes_list.append({
                            "姓名": name,
                            "项目": col,
                            "旧值": old_val if pd.notna(old_val) else "(空)",
                            "新值": val
                        })
            
            if items_updated: 
                logs.append(f"✅ {name}: 更新了 {len(items_updated)} 项")
    
    return merged_df, logs, pd.DataFrame(changes_list)

def smart_clean(df, level):
    df_c = df.copy()
    audit = []
    bar = st.progress(0)
    for i, row in df_c.iterrows():
        name = row.get('姓名', '')
        gender = row.get('性别')
        if '小学' not in level:
            if str(gender) in ['1', '1.0', '男']: 
                for c in ['一分钟仰卧起坐', '仰卧起坐', '800米跑', '800米']:
                    if c in df_c.columns and pd.notna(row[c]): df_c.at[i, c] = None
            elif str(gender) in ['2', '2.0', '女']:
                 for c in ['引体向上', '1000米跑', '1000米']:
                    if c in df_c.columns and pd.notna(row[c]): df_c.at[i, c] = None
        if len(df_c) > 0 and i % 20 == 0: bar.progress((i+1)/len(df_c))
    bar.progress(100)
    time.sleep(0.3)
    bar.empty()
    return df_c, audit

# ==========================================
# 🎮 主程序入口
# ==========================================
render_ui_header()

st.sidebar.markdown("### 🎛️ 控制面板")
api_key_input = st.sidebar.text_input("阿里云 API Key", value=MY_API_KEY, type="password")
st.sidebar.markdown("---")
st.sidebar.markdown("### 🎓 学段选择")
school_level = st.sidebar.selectbox(
    "请选择目标学段",
    ("初中", "高中", "大学", "小学(低年级)", "小学(中年级)", "小学(高年级)"),
    index=0
)

if STANDARDS_DB and school_level in STANDARDS_DB:
    with st.sidebar.expander(f"📊 {school_level}标准预览"):
        try:
            st.caption(f"说明: {STANDARDS_DB[school_level]['说明']}")
            if '男' in STANDARDS_DB[school_level] and '立定跳远' in STANDARDS_DB[school_level]['男']:
                st.caption(f"男生立定跳远及格: {STANDARDS_DB[school_level]['男']['立定跳远']['60']}")
                st.caption(f"女生立定跳远及格: {STANDARDS_DB[school_level]['女']['立定跳远']['60']}")
        except: st.caption("部分预览数据不可用")

# Session State
if 'master_file_obj' not in st.session_state: st.session_state['master_file_obj'] = None
if 'master_filename' not in st.session_state: st.session_state['master_filename'] = ""
if 'master_df' not in st.session_state: st.session_state['master_df'] = None
if 'ocr_df' not in st.session_state: st.session_state['ocr_df'] = None

# === STEP 1: LOAD TEMPLATE ===
st.markdown("#### 📂 第一步：模版总表上传")
master_file = st.file_uploader("请上传包含名单的 Excel 总表", type=['xlsx'], label_visibility="collapsed")

if master_file:
    if st.session_state['master_df'] is None or st.session_state['master_filename'] != master_file.name:
        st.session_state['master_df'] = pd.read_excel(master_file)
        master_file.seek(0)
        st.session_state['master_file_obj'] = io.BytesIO(master_file.read())
        st.session_state['master_filename'] = master_file.name
        st.toast("✅ 模板加载成功！", icon="💿")

if st.session_state['master_df'] is not None:
    st.dataframe(st.session_state['master_df'].head(3), use_container_width=True)
    
    tab1, tab2 = st.tabs(["📸 AI 识图 & 合并", "🚀 数据处理 (清洗+调整)"])

    # === STEP 2: OCR ===
    with tab1:
        st.info("💡 提示：在选择文件窗口中，按住 Ctrl 或 Shift 键可一次性选择多张图片上传。")
        
        # 使用列布局：左侧上传框，右侧按钮
        c1, c2 = st.columns([3, 1], vertical_alignment="bottom")
        
        with c1:
            imgs = st.file_uploader("上传手写成绩单 (支持批量)", type=['jpg', 'png'], accept_multiple_files=True)
            
        with c2:
            # 按钮紧跟上传框，且加大醒目
            start_scan = st.button("✨ 批量 AI 识图", type="primary", use_container_width=True)

        if start_scan and imgs:
            if not api_key_input: st.error("❌ 缺少 API Key")
            else:
                all_data_list = []
                error_logs = []
                
                # 创建进度条
                progress_bar = st.progress(0)
                status_text = st.empty()
                
                # --- 多线程并发加速 ---
                with st.spinner("🚀 AI 引擎全速运转中 (多线程加速)..."):
                    with ThreadPoolExecutor(max_workers=5) as executor:
                        # 提交所有任务
                        future_to_file = {executor.submit(call_qwen_vl_ocr, img, api_key_input): img for img in imgs}
                        
                        completed_count = 0
                        total_files = len(imgs)
                        
                        for future in as_completed(future_to_file):
                            img_file = future_to_file[future]
                            completed_count += 1
                            progress_bar.progress(completed_count / total_files)
                            status_text.text(f"🚀 正在处理第 {completed_count}/{total_files} 张图片...")
                            
                            try:
                                res = future.result()
                                if "Error" in res:
                                    error_logs.append(f"{img_file.name}: {res}")
                                else:
                                    data = json.loads(res)
                                    if isinstance(data, list):
                                        df_temp = pd.DataFrame(data)
                                        all_data_list.append(df_temp)
                                    else:
                                        error_logs.append(f"{img_file.name}: 格式错误")
                            except Exception as e:
                                error_logs.append(f"{img_file.name}: 解析异常 {str(e)}")
                
                # 处理完成
                progress_bar.empty()
                status_text.empty()
                
                if all_data_list:
                    final_ocr_df = pd.concat(all_data_list, ignore_index=True)
                    st.session_state['ocr_df'] = final_ocr_df
                    st.success(f"✅ 成功识别 {len(imgs) - len(error_logs)} 张图片，提取 {len(final_ocr_df)} 条数据")
                
                if error_logs:
                    with st.expander(f"⚠️ {len(error_logs)} 张图片识别失败"):
                        for err in error_logs: st.write(err)

        if st.session_state['ocr_df'] is not None:
            st.markdown("---")
            edited_ocr = st.data_editor(st.session_state['ocr_df'], num_rows="dynamic", use_container_width=True)
            
            # 合并按钮
            if st.button("📥 合并置入总表", type="primary", use_container_width=True):
                new_master, logs, changes_df = merge_ocr_to_master(st.session_state['master_df'], edited_ocr)
                st.session_state['master_df'] = new_master
                
                # 详细合并报告
                if not changes_df.empty:
                    st.markdown("#### 📋 数据合并详情报告")
                    st.dataframe(changes_df, use_container_width=True)
                    st.success(f"🎉 成功更新了 {len(changes_df)} 个数据点！")
                else:
                    st.warning("未检测到有效数据更新 (可能是姓名未匹配或数据为空)")

    # === STEP 3: PROCESS & EXPORT ===
    with tab2:
        col_ctrl, col_info = st.columns([1, 2])
        
        # 实时计算原始优良率
        current_data_rate = calculate_good_rate(st.session_state['master_df'], school_level)
        
        with col_ctrl:
            st.info(f"当前模式：**{school_level}**")
            
            # 显示当前原始状态
            st.metric("📈 原始数据优良率", f"{current_data_rate:.1f}%")
            
            # 双控件联动
            st.write("🎯 设定目标优良率")
            c_slide, c_input = st.columns([2, 1])
            with c_slide:
                rate_slide = st.slider("粗调", 0.0, 100.0, 90.0, 1.0, label_visibility="collapsed")
            with c_input:
                target_rate = st.number_input("精调 (%)", 0.0, 100.0, float(rate_slide), 0.1, label_visibility="collapsed")
                
            run_btn = st.button("⚡ 执行核心算法", type="primary", use_container_width=True)
            
        with col_info:
            if run_btn:
                if STANDARDS_DB is None:
                    st.error("无法运行：找不到评分标准文件 standards.json")
                else:
                    df_clean, _ = smart_clean(st.session_state['master_df'], school_level)
                    df_final = auto_boost(df_clean, target_rate, school_level)
                    st.session_state['master_df'] = df_final
                    
                    # 结果展示
                    achieved_rate = calculate_good_rate(df_final, school_level)
                    st.metric(label="调整后实际优良率 (Good/Excellent Rate)", value=f"{achieved_rate:.1f}%", delta=f"{(achieved_rate - current_data_rate):.1f}%")
                    st.success(f"✅ 基于[{school_level}]国标调整完毕")
        
        st.markdown("---")
        
        if st.session_state['master_df'] is not None:
            excel_data, msg = save_data_keeping_format(st.session_state['master_file_obj'], st.session_state['master_df'])
            if excel_data:
                st.download_button(label="📥 下载最终报表 (XLSX)", data=excel_data, file_name=f"77智能体测_{school_level}_最终.xlsx", mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", use_container_width=True)
            else: st.error(msg)