import streamlit as st
import pandas as pd
import numpy as np
import random
import json
import os
import tempfile
import time
import io
from http import HTTPStatus
from concurrent.futures import ThreadPoolExecutor, as_completed

try:
    import openpyxl
    import dashscope
except ImportError:
    st.error("❌ 严重错误：缺少运行必要的库。请检查 requirements.txt")
    st.stop()

# ==========================================
# 📚 核心配置与标准库管理
# ==========================================
SCORE_COLUMNS_MAP = {
    '50米': ['50米', '50米跑', '五十米', '男生50米', '女生50米', '50m'],
    '立定跳远': ['立定跳远', '跳远', '立定跳远(厘米)', '立定跳远(cm)'],
    '坐位体前屈': ['坐位体前屈', '体前屈', '坐位', '坐位体前屈(厘米)', '坐位体前屈(cm)'],
    '1000米': ['1000米', '1000米跑', '一千米', '男生1000米', '1000m'],
    '800米': ['800米', '800米跑', '八百米', '女生800米', '800m'],
    '引体向上': ['引体向上', '引体', '男生引体向上'],
    '仰卧起坐': ['仰卧起坐', '一分钟仰卧起坐', '女生仰卧起坐'],
    '肺活量': ['肺活量', '肺活量(毫升)', '肺活量(ml)'],
    '身高': ['身高', '身高(厘米)', '身高(cm)'],
    '体重': ['体重', '体重(千克)', '体重(kg)'],
    '跳绳': ['跳绳', '一分钟跳绳', '跳绳(次)'],
    '50x8往返跑': ['50x8', '50米x8', '往返跑', '50*8']
}

# 【修复】重新整理了数据结构，修复了之前的语法错误
DEFAULT_STANDARDS_DATA = {
    "小学(低年级)": {
        "说明": "适用于小学1-2年级",
        "男": {
            "50米": {"60": 12.0, "80": 10.6},
            "坐位体前屈": {"60": -0.6, "80": 6.9},
            "立定跳远": {"60": 117, "80": 140},
            "跳绳": {"60": 17, "80": 80},
            "肺活量": {"60": 680, "80": 1080}
        },
        "女": {
            "50米": {"60": 12.6, "80": 11.2},
            "坐位体前屈": {"60": 1.6, "80": 8.7},
            "立定跳远": {"60": 110, "80": 132},
            "跳绳": {"60": 17, "80": 84},
            "肺活量": {"60": 580, "80": 900}
        }
    },
    "小学(中年级)": {
        "说明": "适用于小学3-4年级",
        "男": {
            "50米": {"60": 10.6, "80": 9.2},
            "坐位体前屈": {"60": 0.5, "80": 8.0},
            "立定跳远": {"60": 138, "80": 161},
            "跳绳": {"60": 40, "80": 95},
            "仰卧起坐": {"60": 23, "80": 39},
            "肺活量": {"60": 1100, "80": 1700}
        },
        "女": {
            "50米": {"60": 10.8, "80": 9.5},
            "坐位体前屈": {"60": 2.8, "80": 9.8},
            "立定跳远": {"60": 130, "80": 149},
            "跳绳": {"60": 42, "80": 99},
            "仰卧起坐": {"60": 21, "80": 37},
            "肺活量": {"60": 950, "80": 1350}
        }
    },
    "小学(高年级)": {
        "说明": "适用于小学5-6年级",
        "男": {
            "50米": {"60": 9.8, "80": 8.6},
            "坐位体前屈": {"60": 2.0, "80": 9.0},
            "立定跳远": {"60": 161, "80": 193},
            "跳绳": {"60": 65, "80": 126},
            "仰卧起坐": {"60": 28, "80": 42},
            "50x8往返跑": {"60": 138, "80": 122},
            "肺活量": {"60": 1400, "80": 2200}
        },
        "女": {
            "50米": {"60": 10.2, "80": 9.1},
            "坐位体前屈": {"60": 4.5, "80": 10.5},
            "立定跳远": {"60": 139, "80": 164},
            "跳绳": {"60": 67, "80": 129},
            "仰卧起坐": {"60": 27, "80": 41},
            "50x8往返跑": {"60": 144, "80": 128},
            "肺活量": {"60": 1150, "80": 1650}
        }
    },
    "初中": {
        "说明": "适用于初中7-9年级",
        "男": {
            "50米": {"60": 9.1, "80": 7.9},
            "坐位体前屈": {"60": 3.7, "80": 12.8},
            "立定跳远": {"60": 185, "80": 225},
            "引体向上": {"60": 6, "80": 11},
            "1000米": {"60": 275, "80": 245},
            "肺活量": {"60": 2300, "80": 3400}
        },
        "女": {
            "50米": {"60": 9.5, "80": 8.5},
            "坐位体前屈": {"60": 6.0, "80": 14.8},
            "立定跳远": {"60": 146, "80": 176},
            "仰卧起坐": {"60": 28, "80": 40},
            "800米": {"60": 255, "80": 225},
            "肺活量": {"60": 1800, "80": 2500}
        }
    },
    "高中": {
        "说明": "适用于高中10-12年级",
        "男": {
            "50米": {"60": 9.0, "80": 7.8},
            "坐位体前屈": {"60": 5.5, "80": 14.9},
            "立定跳远": {"60": 205, "80": 240},
            "引体向上": {"60": 9, "80": 14},
            "1000米": {"60": 265, "80": 235},
            "肺活量": {"60": 2700, "80": 3800}
        },
        "女": {
            "50米": {"60": 9.4, "80": 8.7},
            "坐位体前屈": {"60": 6.8, "80": 16.5},
            "立定跳远": {"60": 158, "80": 182},
            "仰卧起坐": {"60": 28, "80": 40},
            "800米": {"60": 270, "80": 240},
            "肺活量": {"60": 1850, "80": 2650}
        }
    },
    "大学": {
        "说明": "适用于大学1-4年级",
        "男": {
            "50米": {"60": 9.1, "80": 7.9},
            "坐位体前屈": {"60": 3.7, "80": 14.7},
            "立定跳远": {"60": 208, "80": 248},
            "引体向上": {"60": 10, "80": 15},
            "1000米": {"60": 272, "80": 237},
            "肺活量": {"60": 3100, "80": 4300}
        },
        "女": {
            "50米": {"60": 10.3, "80": 8.9},
            "坐位体前屈": {"60": 6.0, "80": 16.9},
            "立定跳远": {"60": 151, "80": 178},
            "仰卧起坐": {"60": 26, "80": 42},
            "800米": {"60": 274, "80": 244},
            "肺活量": {"60": 2000, "80": 3000}
        }
    }
}

@st.cache_data(ttl=3600)
def load_standards():
    file_path = 'standards.json'
    # 如果文件不存在，写入默认数据
    if not os.path.exists(file_path):
        try:
            with open(file_path, 'w', encoding='utf-8') as f:
                json.dump(DEFAULT_STANDARDS_DATA, f, ensure_ascii=False, indent=2)
        except: return None
    # 尝试读取
    try:
        with open(file_path, 'r', encoding='utf-8') as f: return json.load(f)
    except: return None

STANDARDS_DB = load_standards()

# ==========================================
# 🧠 核心算法
# ==========================================
def calculate_good_rate(df, level):
    if not STANDARDS_DB or level not in STANDARDS_DB: return 0.0
    std_male, std_female = STANDARDS_DB[level]["男"], STANDARDS_DB[level]["女"]
    total_items, total_good = 0, 0
    if len(df) == 0: return 0.0
    for _, row in df.iterrows():
        gender = row.get('性别')
        current_std = std_male if str(gender) in ['1','1.0','男'] else std_female
        for col in df.columns:
            key = next((k for k in current_std.keys() if k in col), None)
            if key and pd.notna(row[col]):
                try:
                    val, good_cutoff = float(row[col]), float(current_std[key]['80'])
                    is_running = key in ['50米', '800米', '1000米', '50x8往返跑', '50米跑', '800米跑', '1000米跑']
                    if (is_running and val <= good_cutoff) or (not is_running and val >= good_cutoff):
                        total_good += 1
                    total_items += 1
                except: pass
    return (total_good / total_items) * 100 if total_items > 0 else 0.0

def auto_boost(df, target_rate, level):
    df_a = df.copy()
    boost_changes = []
    if not STANDARDS_DB or level not in STANDARDS_DB: return df_a, pd.DataFrame()
    
    std_male, std_female = STANDARDS_DB[level]["男"], STANDARDS_DB[level]["女"]
    not_good_candidates = [] 
    total_valid_items, current_good_items = 0, 0
    
    for idx, row in df_a.iterrows():
        gender = row.get('性别')
        current_std = std_male if str(gender) in ['1','1.0','男'] else std_female
        for col in df_a.columns:
            key = next((k for k in current_std.keys() if k in col), None)
            if key and pd.notna(row[col]):
                try:
                    val = float(row[col])
                    good_cutoff = float(current_std[key]['80'])
                    is_running = key in ['50米', '800米', '1000米', '50x8往返跑', '50米跑', '800米跑', '1000米跑']
                    is_good = (is_running and val <= good_cutoff) or (not is_running and val >= good_cutoff)
                    total_valid_items += 1
                    if is_good: current_good_items += 1
                    else:
                        distance = (val - good_cutoff) if is_running else (good_cutoff - val)
                        not_good_candidates.append({
                            'idx': idx, 'col': col, 'val': val,
                            'threshold': good_cutoff, 'is_running': is_running,
                            'distance': distance
                        })
                except: continue
                
    if total_valid_items == 0: return df_a, pd.DataFrame()
    target_good_count = int(total_valid_items * (target_rate / 100))
    needed = target_good_count - current_good_items
    
    with st.status(f"🚀 智能调整算法执行中... (目标: {target_rate}%)", expanded=True) as status:
        if needed <= 0:
            st.success("✅ 已达标，无需调整。")
            time.sleep(0.5)
            status.update(state="complete", expanded=False)
            return df_a, pd.DataFrame()
        
        st.write(f"⚡ 正在优先优化 {needed} 个最接近目标的数据点...")
        not_good_candidates.sort(key=lambda x: x['distance'])
        cells_to_boost = not_good_candidates[:needed]
        prog = st.progress(0)
        
        for i, cell in enumerate(cells_to_boost):
            row_idx, col_name = cell['idx'], cell['col']
            threshold, is_run = cell['threshold'], cell['is_running']
            old_val = df_a.at[row_idx, col_name]
            student_name = df_a.at[row_idx, '姓名'] if '姓名' in df_a.columns else f"行{row_idx+1}"

            if is_run:
                boost_amount = random.uniform(0.01, 0.2)
                new_val = round(max(0.1, threshold - boost_amount), 2)
            elif any(x in col_name for x in ['跳绳','仰卧','引体']):
                new_val = int(threshold + random.randint(0, 2))
            elif any(x in col_name for x in ['跳远','身高','体重']):
                new_val = int(threshold + random.randint(0, 3))
            elif '肺活量' in col_name:
                new_val = int(threshold + random.randint(10, 100))
            else:
                new_val = round(threshold + random.uniform(0.1, 1.0), 1)
                
            df_a.at[row_idx, col_name] = new_val
            boost_changes.append({"姓名": student_name, "调整项目": col_name, "调整前(旧值)": old_val, "调整后(新值)": new_val})
            if i % (max(1, len(cells_to_boost)//10)) == 0: prog.progress((i+1)/len(cells_to_boost))
        prog.progress(100)
        status.update(label=f"✅ 调整完毕！", state="complete", expanded=False)
        
    return df_a, pd.DataFrame(boost_changes)

# ==========================================
# 🛠️ 工具函数
# ==========================================
def save_data_keeping_format(orig_file_bytes_io, processed_df):
    try:
        wb = openpyxl.load_workbook(orig_file_bytes_io)
        ws = wb.active
        header_map = {str(cell.value).strip(): i for i, cell in enumerate(ws[1], 1) if cell.value}
        name_col_idx = header_map.get('姓名')
        if not name_col_idx: return None, "❌ 错误：在模板中未找到【姓名】列。"
        name_row_map = {}
        for i, row in enumerate(ws.iter_rows(min_row=2), 2):
            cell_val = row[name_col_idx-1].value
            if cell_val: name_row_map[str(cell_val).strip()] = i
        for _, row in processed_df.iterrows():
            name = str(row.get('姓名','')).strip()
            target_row = name_row_map.get(name)
            if target_row:
                for col_name, val in row.items():
                    target_col = header_map.get(col_name)
                    if target_col and pd.notna(val): ws.cell(target_row, target_col).value = val
        out = io.BytesIO(); wb.save(out); out.seek(0)
        return out, "✅ 成功"
    except Exception as e: return None, f"导出失败: {str(e)}"

def call_qwen_vl_ocr(img_file, api_key, filename):
    dashscope.api_key = api_key
    prompt = "你是一个专业的体测数据录入员。请分析这张体测成绩单图片。任务：提取表格中的姓名、性别、班级，以及所有体育项目的成绩。重要规则：1. **姓名必须极其准确**，这是匹配的关键。2. 如果有“立定跳远”项目，且单位是“米”（例如 2.3），请务必转换为“厘米”（例如 230）。3. 最终输出必须是一个纯粹的 JSON 格式列表（List of Dicts），不要包含任何Markdown标记（如 ```json ... ```）或其他解释文字。例如：[{\"姓名\": \"张三\", \"性别\": \"男\", \"50米\": \"7.5\", \"立定跳远\": \"230\"}, ...]"
    local_path = None
    try:
        img_file.seek(0)
        with tempfile.NamedTemporaryFile(delete=False, suffix=img_file.name) as t:
            t.write(img_file.read()); local_path = t.name
        messages = [{"role": "user", "content": [{"image": f"file://{local_path}"}, {"text": prompt}]}]
        resp = dashscope.MultiModalConversation.call(model='qwen-vl-max', messages=messages)
        os.remove(local_path)
        if resp.status_code == HTTPStatus.OK:
            content = resp.output.choices[0].message.content[0]['text'].replace("```json", "").replace("```", "").strip()
            data = json.loads(content)
            if isinstance(data, list):
                for item in data: item['来源图片'] = filename
            return data
        else: return f"Error: API 返回状态码 {resp.code} - {resp.message}"
    except Exception as e:
        if local_path and os.path.exists(local_path): os.remove(local_path)
        return f"Error: {str(e)}"

def generate_class_report(df, level, rate, api_key):
    dashscope.api_key = api_key
    stats_text = f"当前学段：{level}，班级人数：{len(df)}人。经过调整后，优良率达到了：{rate:.1f}%。"
    prompt = f"""
    你是一位资深的体育特级教师。请根据以下班级体测概况，写一份简短的《体育教学分析与改进建议》。
    数据概况：{stats_text}
    要求：1. 语气专业、鼓励性强。2. 第一段点评整体情况。3. 第二段给出针对性的训练建议。4. 300字以内。
    """
    try:
        resp = dashscope.Generation.call(model='qwen-turbo', messages=[{'role': 'user', 'content': prompt}])
        if resp.status_code == HTTPStatus.OK: return resp.output.text
        else: return "无法生成报告，请稍后再试。"
    except: return "网络连接异常，无法生成报告。"

def merge_ocr_to_master(master_df, ocr_df):
    merged_df = master_df.copy()
    logs, changes = [], []
    ocr_cols = ocr_df.columns
    col_map = {}
    for std_col, alias_list in SCORE_COLUMNS_MAP.items():
        for alias in alias_list:
            if alias in ocr_cols: col_map[alias] = std_col; break
    ocr_df = ocr_df.rename(columns=col_map)
    with st.spinner("🔄 正在进行智能匹配与数据合并..."):
        for _, row in ocr_df.iterrows():
            name = str(row.get('姓名','')).strip()
            if not name: continue
            mask = merged_df['姓名'].astype(str).str.strip() == name
            if mask.any():
                target_idx = merged_df[mask].index[0]
                updated_items = []
                for col in row.index:
                    if col != '来源图片' and col in SCORE_COLUMNS_MAP and col in merged_df.columns and pd.notna(row[col]):
                        old_val = merged_df.at[target_idx, col]
                        merged_df.at[target_idx, col] = row[col]
                        updated_items.append(col)
                        changes.append({"姓名": name, "项目": col, "旧值": old_val if pd.notna(old_val) else "(空)", "新值": row[col], "来源图片": row.get('来源图片', '未知')})
                if updated_items: logs.append(f"✅ {name}: 成功更新 {len(updated_items)} 个项目")
    return merged_df, logs, pd.DataFrame(changes)

def smart_clean(df, level):
    df_c = df.copy()
    with st.spinner("🧹 正在根据学段和性别清洗无关数据..."):
        for i, row in df_c.iterrows():
            gender = str(row.get('性别'))
            if '小学' not in level:
                if gender in ['1', '1.0', '男']:
                    for c in ['一分钟仰卧起坐', '仰卧起坐', '800米跑', '800米']:
                        if c in df_c.columns: df_c.at[i, c] = None
                elif gender in ['2', '2.0', '女']:
                    for c in ['引体向上', '1000米跑', '1000米']:
                        if c in df_c.columns: df_c.at[i, c] = None
    return df_c, []

def validate_data_ranges(df):
    anomalies = []
    bounds = {
        '50米': (5.0, 25.0), '50米跑': (5.0, 25.0), '800米': (100, 600), '800米跑': (100, 600),
        '1000米': (120, 800), '1000米跑': (120, 800), '立定跳远': (50, 400), '跳远': (50, 400),
        '身高': (80, 250), '体重': (15, 200), '肺活量': (300, 10000)
    }
    for idx, row in df.iterrows():
        name = row.get('姓名', f'行{idx+1}')
        for col, val in row.items():
            if col in bounds and pd.notna(val):
                try:
                    v = float(val)
                    lower, upper = bounds[col]
                    if not (lower <= v <= upper):
                        anomalies.append({"姓名": name, "项目": col, "异常值": v, "异常原因": f"超出合理范围 {lower}-{upper}"})
                except:
                    anomalies.append({"姓名": name, "项目": col, "异常值": str(val), "异常原因": "格式错误(非数字)"})
    return pd.DataFrame(anomalies)